#!/usr/bin/env python3
"""
Consolida a planilha PRICO do cliente (Material -> Produto -> SKU) para carga no MongoDB.

Uso:
    pip install python-calamine openpyxl
    python tools/prico/consolidar_prico.py <planilha.xlsx> [--repo .] [--report relatorio.xlsx]

Entradas:
    - Planilha do cliente (usa somente 37_CATALOGO_APTO com PROTOCOLO PARA INSERÇÃO = OK,
      conforme 29_DIRETRIZES_IMPLANTACAO), mais 03/28 (contrato de materiais e variações)
      e 10/11 (datas de consulta e origem dos mesmos registros).
    - Catálogo atual do backend: src/main/resources/catalog/prico_variacoes.ndjson.gz.b64

Saídas (no repositório):
    - src/main/resources/catalog/prico_variacoes.ndjson.gz.b64   catálogo v2 (materiais/variações/opções)
    - src/main/resources/catalog/prico_produtos.ndjson.gz          um produto (com SKUs embutidos) por linha
    - src/main/resources/catalog/prico_produtos.manifest.json      lote, hash da origem e contagens
    - relatório de revisão (.xlsx) com tudo que foi tratado, excluído ou precisa voltar ao cliente

O script é determinístico: a mesma planilha gera os mesmos arquivos. Nada é inventado; valores
ausentes ficam ausentes e toda normalização preserva o valor original em campo próprio.
"""
from __future__ import annotations

import argparse
import base64
import datetime as dt
import gzip
import hashlib
import io
import json
import re
import sys
import unicodedata
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

from python_calamine import CalamineWorkbook

CATALOG_RESOURCE = "src/main/resources/catalog/prico_variacoes.ndjson.gz.b64"
PRODUCTS_RESOURCE = "src/main/resources/catalog/prico_produtos.ndjson.gz"
MANIFEST_RESOURCE = "src/main/resources/catalog/prico_produtos.manifest.json"

DEFAULT_MATERIAL_STATUS = "EM_REVISÃO"
DEFAULT_MATERIAL_OBSERVATION = "Material canônico, independente de marca/fabricante."
CONTRACT_ACTIVE = "ATIVO"
CONTRACT_OUT = "FORA_DO_CONTRATO"

# --------------------------------------------------------------------------------------------
# Leitura
# --------------------------------------------------------------------------------------------

def cell(value):
    """Converte a célula do calamine em texto estável (None para vazio)."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    text = str(value)
    return text if text.strip() else None


def read_sheet(workbook, name):
    rows = workbook.get_sheet_by_name(name).to_python(skip_empty_area=False)
    header = [str(h).strip() for h in rows[0]]
    raw_types = []
    out = []
    for raw in rows[1:]:
        if all(v in (None, "") for v in raw):
            continue
        record = {header[i]: cell(v) for i, v in enumerate(raw) if i < len(header)}
        record["__types__"] = {header[i]: type(v).__name__ for i, v in enumerate(raw) if i < len(header)}
        out.append(record)
    return out


def read_bundled_catalog(repo: Path):
    encoded = (repo / CATALOG_RESOURCE).read_bytes()
    data = gzip.decompress(base64.b64decode(encoded))
    return [json.loads(line) for line in data.decode("utf-8").splitlines() if line.strip()]


# --------------------------------------------------------------------------------------------
# Normalizações
# --------------------------------------------------------------------------------------------

def strip_accents(text):
    return unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode()


def key(text):
    return re.sub(r"\s+", " ", strip_accents(str(text or "")).lower()).strip()


def clean_text(text):
    if text is None:
        return None
    text = str(text).replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    text = "\n".join(line for line in lines if line)
    return text or None


def code_sort_key(code):
    parts = re.split(r"[.]", code)
    out = []
    for part in parts:
        m = re.fullmatch(r"([A-Z]*)(\d+)", part)
        out.append((m.group(1), int(m.group(2))) if m else (part, 0))
    return out


def iso_date(value):
    if not value:
        return None
    if re.fullmatch(r"\d{5}", value):  # serial do Excel
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(value))).isoformat()
    m = re.match(r"(\d{4}-\d{2}-\d{2})", value)
    return m.group(1) if m else None


def gtin_valid(gtin):
    digits = re.sub(r"\D", "", gtin)
    if len(digits) not in (8, 12, 13, 14) or digits != gtin:
        return False
    nums = [int(d) for d in digits]
    body = nums[:-1][::-1]
    total = sum(n * (3 if i % 2 == 0 else 1) for i, n in enumerate(body))
    return (10 - total % 10) % 10 == nums[-1]


def format_number(value):
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
        return str(int(value))
    text = f"{value:.6f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


NUMBER = r"\d+(?:[.,]\d+)?"
DIM_SPLIT = re.compile(r"\s*[x×X*]\s*")


def parse_dimensions(text):
    parts = DIM_SPLIT.split(text.strip())
    if not parts or not all(re.fullmatch(NUMBER, p) for p in parts):
        return None, text
    numbers = [float(p.replace(",", ".")) for p in parts]
    display = " x ".join(p for p in parts)
    return numbers, display


# Unidade comercial: o cliente mistura unidade e conteúdo ("Caixa com 16 placas").
# A unidade base vai para commercialUnit e o texto original (quando difere) para commercialUnitDetail.
UNIT_SYNONYMS = {
    "un": "un", "und": "un", "unid": "un", "unidade": "un", "u": "un",
    "pc": "peça", "pç": "peça", "peca": "peça", "peça": "peça",
    "cx": "caixa", "caixa": "caixa",
    "sc": "saco", "saco": "saco",
    "pct": "pacote", "pacote": "pacote", "pack": "pacote",
    "m": "m", "metro": "m",
    "t": "t", "ton": "t", "tonelada": "t",
    "kg": "kg", "g": "g", "l": "L", "ml": "mL",
    "m²": "m²", "m2": "m²", "m³": "m³", "m3": "m³",
    "pallet": "palete", "palete": "palete",
}
UNIT_NOUNS = {
    "rolo", "conjunto", "embalagem", "kit", "placa", "carretel", "par", "cento", "chapa", "barra",
    "bobina", "galão", "balde", "tambor", "jogo", "cartucho", "pé", "bombona", "módulo", "lata",
    "painel", "contêiner", "porta", "blister", "unipack", "barrica", "porção", "régua", "folha",
    "manta", "perfil", "pote", "tijolo", "telha", "equipamento", "lâmina", "recipiente", "aerossol",
    "amarrado", "bandeja", "bisnaga", "cartela", "cilindro", "escada", "fita", "grelha", "ibc",
    "ladrilho", "lance", "latinha", "milheiro", "quarto", "saqueta", "sistema", "bloco", "frasco",
    "tubo", "big bag", "kit de fornecimento", "frasco/tubo", "saco/embalagem",
}


def normalize_unit(raw):
    """Retorna (unidade_base, detalhe_ou_None, padronizada?)."""
    if raw is None:
        return None, None, True
    text = clean_text(raw)
    low = text.lower()
    if low in UNIT_SYNONYMS:
        return UNIT_SYNONYMS[low], None if UNIT_SYNONYMS[low] == text else text, True
    if low in UNIT_NOUNS:
        return low, None if low == text else text, True
    m = re.match(r"([a-zà-ú²³/]+(?: bag)?)\b", low)
    head = m.group(1) if m else None
    if head in UNIT_SYNONYMS:
        return UNIT_SYNONYMS[head], text, True
    if head in UNIT_NOUNS:
        return head, text, True
    if "(peça)" in low:
        return "peça", text, True
    return None, text, False


NAME_STOPWORDS = {"de", "da", "do", "das", "dos", "para", "com", "por", "em", "e", "ou", "nominal", "nominais",
                  "total", "maxima", "maximo", "minima", "minimo"}
TYPE_GROUPS = [
    {"DIMENSAO", "DIMENSOES", "COMPRIMENTO", "LARGURA", "ALTURA", "ESPESSURA", "DIAMETRO", "BITOLA", "TAMANHO", "AREA", "CONEXAO"},
    {"EMBALAGEM", "CONTEUDO", "MASSA", "VOLUME", "QUANTIDADE", "GRAMATURA", "UNIDADE"},
    {"COR", "ACABAMENTO", "TEXTO"},
    {"POTENCIA", "CAPACIDADE", "ENERGIA", "VAZAO"},
    {"CLASSE", "RESISTENCIA", "DENSIDADE", "PRESSAO"},
]


def name_stems(text):
    """Radicais (5 letras) do nome, sem palavras de ligação: 'Dimensão nominal' ~ 'Dimensões nominais'."""
    return {t[:5] for t in re.split(r"[^a-z0-9]+", key(text)) if len(t) >= 3 and t not in NAME_STOPWORDS}


def same_type_group(a, b):
    a, b = key(a).upper(), key(b).upper()
    return a == b or any(a in group and b in group for group in TYPE_GROUPS)


def canonical(values):
    """Escolhe a grafia mais frequente dentro de um grupo (empate: evita CAIXA ALTA)."""
    counts = Counter(values)
    return sorted(counts, key=lambda v: (-counts[v], v.isupper(), v))[0]


def option_keys(*parts):
    text = "".join(str(p) for p in parts if p not in (None, ""))
    text = unicodedata.normalize("NFKC", text).lower()
    return re.sub(r"\s+", "", text).replace(",", ".")


# --------------------------------------------------------------------------------------------
# Controles do próprio cliente (registro com prefixo LEN)
# --------------------------------------------------------------------------------------------

REGISTER_COLUMNS_37 = [
    "COD MATERIAL", "COD PRODUTO", "PRODUTO", "MARCA", "FABRICANTE", "LINHA/MODELO", "COD SKU PRICO",
    "SKU FABRICANTE", "GTIN/EAN", "UNIDADE COMERCIAL", "APRESENTAÇÃO", "FONTE PRODUTO", "FONTE SKU",
    "VARIAÇÕES JSON", "PENDÊNCIAS DE CAMPOS", "COBERTURA DOCUMENTAL", "ORIGEM UNIÃO", "OBSERVAÇÃO",
    "MODALIDADE COMERCIAL", "CHAVE IDENTIDADE COMERCIAL", "CONTEÚDOS VALIDADOS",
]


def parse_register(text):
    fields, i = [], 0
    while i < len(text):
        j = text.index(":", i)
        size = int(text[i:j])
        fields.append(text[j + 2:j + 1 + size])
        i = j + 1 + size
    return fields


def register_problem(row):
    if row.get("PROTOCOLO PARA INSERÇÃO") != "OK":
        return "PROTOCOLO_DIFERENTE_DE_OK"
    if row.get("CONTROLE CATÁLOGO") != "OK":
        return "CONTROLE_CATALOGO_DIFERENTE_DE_OK"
    if row.get("PENDÊNCIAS DA AUDITORIA"):
        return "PENDENCIA_DE_AUDITORIA"
    if row.get("REGISTRO ATUAL") != row.get("REGISTRO VALIDADO"):
        return "REGISTRO_ATUAL_DIVERGE_DO_VALIDADO"
    try:
        fields = parse_register(row["REGISTRO ATUAL"])
    except Exception:
        return "REGISTRO_ILEGIVEL"
    for idx, column in enumerate(REGISTER_COLUMNS_37):
        if idx >= len(fields) or fields[idx] != (row.get(column) or ""):
            return f"CELULA_DIFERE_DO_REGISTRO:{column} (recalcular a planilha)"
    return None


# --------------------------------------------------------------------------------------------
# Catálogo v2 (materiais + variações + opções)
# --------------------------------------------------------------------------------------------

def build_catalog(v1_rows, contracts, coverage, report):
    v1_materials = OrderedDict()
    v1_variations = OrderedDict()
    v1_options = defaultdict(list)
    for r in v1_rows:
        v1_materials.setdefault(r["Cod. MT"], r)
        v1_variations.setdefault(r["Cod. VR"], r)
        if r.get("Cod. OP"):
            v1_options[r["Cod. VR"]].append(r)
    segments = {}
    families = {}
    for r in v1_rows:
        segments[int(r["Cod. SG"])] = r["SEGMENTO"]
        families[r["Cod. FM"]] = r["FAMÍLIA"]

    cov = {r["COD MATERIAL"]: r for r in coverage}
    by_material = defaultdict(list)
    for c in contracts:
        by_material[c["COD MATERIAL"]].append(c)

    rows = []
    materials_out = {}
    for code in sorted(set(by_material) | set(v1_materials), key=code_sort_key):
        old = v1_materials.get(code)
        new = cov.get(code)
        seg_code = int(code.split(".")[0])
        fam_code = code.rsplit(".", 1)[0]
        if seg_code not in segments:
            raise SystemExit(f"Segmento {seg_code} sem nome no catálogo atual ({code}).")
        if new:
            family_name = clean_text(new["FAMÍLIA"])
            material_name = clean_text(new["MATERIAL"])
        else:
            family_name = old["FAMÍLIA"]
            material_name = old["MATERIAL"]
        base = {
            "Cod. SG": seg_code, "SEGMENTO": segments[seg_code],
            "Cod. FM": fam_code, "FAMÍLIA": family_name,
            "Cod. MT": code, "MATERIAL": material_name,
            "STATUS_MT": old["STATUS_MT"] if old else DEFAULT_MATERIAL_STATUS,
            "OBSERVAÇÃO_MT": old["OBSERVAÇÃO_MT"] if old else DEFAULT_MATERIAL_OBSERVATION,
        }
        if old is None:
            report["catalog_new_material"].append([code, family_name, material_name,
                                                   "NOVA" if fam_code not in families else "existente"])
        elif new is None:
            report["catalog_material_not_in_contract"].append([code, old["FAMÍLIA"], old["MATERIAL"]])
        else:
            if key(old["MATERIAL"]) != key(material_name):
                report["catalog_material_renamed"].append([code, old["MATERIAL"], material_name])
            if key(old["FAMÍLIA"]) != key(family_name):
                report["catalog_family_renamed"].append([code, fam_code, old["FAMÍLIA"], family_name])
        materials_out[code] = {"material": base, "variations": {}}

        contract_by_code = {c["COD VARIAÇÃO"]: c for c in by_material.get(code, [])}
        old_var_codes = [v for v, r in v1_variations.items() if r["Cod. MT"] == code]
        for vcode in sorted(set(contract_by_code) | set(old_var_codes), key=code_sort_key):
            ct = contract_by_code.get(vcode)
            ov = v1_variations.get(vcode)
            order = int(vcode.rsplit(".V", 1)[1])
            if ct:
                units = [u.strip() for u in re.split(r"\|", ct.get("UNIDADES ACEITAS") or "") if u.strip()]
                var = {
                    "Cod. VR": vcode, "VARIAÇÃO": clean_text(ct["ATRIBUTO"]),
                    "TIPO": clean_text(ct.get("TIPO LÓGICO")) or (ov or {}).get("TIPO"),
                    "OBRIGATÓRIA": ct["OBRIGATORIEDADE"], "ORDEM_VR": order,
                    "FORMATO": ct.get("FORMATO"), "UNIDADES_ACEITAS": units,
                    "REGRA_IDENTIDADE": clean_text(ct.get("REGRA DE IDENTIDADE")),
                    "STATUS_CONTRATO": CONTRACT_ACTIVE,
                }
                if ov is None:
                    report["catalog_new_variation"].append([vcode, material_name, var["VARIAÇÃO"], var["FORMATO"]])
                elif key(ov["VARIAÇÃO"]) != key(var["VARIAÇÃO"]):
                    if not (name_stems(ov["VARIAÇÃO"]) & name_stems(var["VARIAÇÃO"])) and not same_type_group(ov.get("TIPO"), var["TIPO"]):
                        meaning = "REVISAR: código reaproveitado com outro significado"
                    else:
                        meaning = "renomeada (mesmo significado)"
                    report["catalog_variation_renamed"].append(
                        [vcode, ov["VARIAÇÃO"], var["VARIAÇÃO"], ov.get("TIPO"), var["TIPO"],
                         len(v1_options.get(vcode, [])), meaning])
            else:
                var = {
                    "Cod. VR": vcode, "VARIAÇÃO": ov["VARIAÇÃO"], "TIPO": ov["TIPO"],
                    "OBRIGATÓRIA": ov["OBRIGATÓRIA"], "ORDEM_VR": order,
                    "FORMATO": None, "UNIDADES_ACEITAS": [], "REGRA_IDENTIDADE": None,
                    "STATUS_CONTRATO": CONTRACT_OUT,
                }
                report["catalog_variation_not_in_contract"].append(
                    [vcode, ov["MATERIAL"], ov["VARIAÇÃO"], len(v1_options.get(vcode, []))])
            options = v1_options.get(vcode, [])
            materials_out[code]["variations"][vcode] = {"variation": var, "options": options}
            if not options:
                rows.append({**base, **var, "Nº OP": None, "Cod. OP": None, "OPÇÃO": None,
                             "SÍMBOLO / UNIDADE": None, "ORDEM_OP": None, "STATUS_OP": None})
            for o in options:
                rows.append({**base, **var, "Nº OP": o.get("Nº OP"), "Cod. OP": o["Cod. OP"],
                             "OPÇÃO": o["OPÇÃO"], "SÍMBOLO / UNIDADE": o.get("SÍMBOLO / UNIDADE"),
                             "ORDEM_OP": o.get("ORDEM_OP"), "STATUS_OP": o.get("STATUS_OP")})
    return rows, materials_out


# --------------------------------------------------------------------------------------------
# Produtos e SKUs
# --------------------------------------------------------------------------------------------

def canonical_product_code(code):
    material, seq = code.rsplit(".P", 1)
    return f"{material}.P{int(seq):04d}" if len(seq) < 4 else code


def build_products(apto, products_base, skus_base, catalog, report, batch_id):
    base_p = {r["COD PRODUTO"]: r for r in products_base}
    base_s = {r["COD SKU PRICO"]: r for r in skus_base}

    accepted = []
    for row in apto:
        problem = register_problem(row)
        if problem:
            report["excluded_rows"].append([row.get("COD SKU PRICO"), row.get("COD PRODUTO"), problem])
            continue
        if row["COD MATERIAL"] not in catalog:
            report["excluded_rows"].append([row["COD SKU PRICO"], row["COD PRODUTO"], "MATERIAL_INEXISTENTE_NO_CATALOGO_V2"])
            continue
        accepted.append(row)

    # Grafias canônicas de marca/fabricante (aplicadas no lote inteiro).
    brand_groups = defaultdict(list)
    maker_groups = defaultdict(list)
    for row in accepted:
        if row.get("MARCA"):
            brand_groups[key(row["MARCA"])].append(clean_text(row["MARCA"]))
        if row.get("FABRICANTE"):
            maker_groups[key(row["FABRICANTE"])].append(clean_text(row["FABRICANTE"]))
    brand_canon = {k: canonical(v) for k, v in brand_groups.items()}
    maker_canon = {k: canonical(v) for k, v in maker_groups.items()}
    for label, groups, canon in (("MARCA", brand_groups, brand_canon), ("FABRICANTE", maker_groups, maker_canon)):
        for k, values in groups.items():
            variants = sorted(set(values))
            if len(variants) > 1:
                report["brand_variants"].append([label, canon[k], " | ".join(variants), len(values)])

    products = OrderedDict()
    sku_seq_used = defaultdict(set)
    pending_bare = []
    for row in accepted:
        source_pcode = row["COD PRODUTO"]
        pcode = canonical_product_code(source_pcode)
        material = row["COD MATERIAL"]
        if not source_pcode.startswith(material + ".P"):
            raise SystemExit(f"Produto {source_pcode} não pertence ao material {material}.")
        bp = base_p.get(source_pcode, {})
        product = products.get(pcode)
        if product is None:
            product = OrderedDict(
                productCode=pcode,
                sourceCode=source_pcode if pcode != source_pcode else None,
                materialCode=material,
                familyCode=material.rsplit(".", 1)[0],
                segmentCode=material.split(".")[0],
                name=clean_text(row["PRODUTO"]),
                brand=brand_canon.get(key(row["MARCA"])) if row.get("MARCA") else None,
                manufacturer=maker_canon.get(key(row["FABRICANTE"])) if row.get("FABRICANTE") else None,
                model=clean_text(row.get("LINHA/MODELO")),
                fingerprint=bp.get("FINGERPRINT"),
                sourceUrl=clean_text(row.get("FONTE PRODUTO")),
                documentalCoverage=row.get("COBERTURA DOCUMENTAL"),
                consultedAt=iso_date(bp.get("DATA CONSULTA")),
                origin=OrderedDict(union=row.get("ORIGEM UNIÃO"), file=bp.get("ARQUIVO ORIGEM"),
                                   row=int(bp["LINHA ORIGEM"]) if (bp.get("LINHA ORIGEM") or "").isdigit() else None),
                skus=[],
                importBatch=batch_id,
            )
            products[pcode] = product
            if product["brand"] and product["brand"] != clean_text(row["MARCA"]):
                report["brand_applied"].append([pcode, row["MARCA"], product["brand"]])

        source_scode = row["COD SKU PRICO"]
        m = re.fullmatch(re.escape(source_pcode) + r"\.S(\d+)", source_scode)
        if m:
            scode = f"{pcode}.S{int(m.group(1)):04d}"
            sku_seq_used[pcode].add(int(m.group(1)))
        else:
            scode = None  # código solto (ex.: S00001): numerado depois, sem colidir
            pending_bare.append((pcode, source_scode))
        bs = base_s.get(source_scode, {})

        unit, unit_detail, unit_ok = normalize_unit(row.get("UNIDADE COMERCIAL"))
        if not unit_ok:
            report["unit_not_standard"].append([source_scode, row.get("UNIDADE COMERCIAL")])
        manufacturer_sku = clean_text(row.get("SKU FABRICANTE"))
        if manufacturer_sku and row["__types__"].get("SKU FABRICANTE") == "float":
            report["numeric_cells"].append([source_scode, "SKU FABRICANTE", manufacturer_sku,
                                            "célula numérica no Excel: confirmar zeros/decimais"])
        gtin = clean_text(row.get("GTIN/EAN"))
        if gtin and not gtin_valid(gtin):
            report["gtin_invalid"].append([source_scode, gtin])
            gtin = None

        values = []
        for item in json.loads(row["VARIAÇÕES JSON"]):
            values.append(build_value(item, material, catalog, source_scode, row.get("FONTE SKU"), report))

        sku = OrderedDict(
            skuCode=scode,
            sourceCode=source_scode,
            manufacturerSku=manufacturer_sku,
            gtin=gtin,
            commercialUnit=unit,
            commercialUnitDetail=unit_detail,
            presentation=clean_text(row.get("APRESENTAÇÃO")),
            sourceUrl=clean_text(row.get("FONTE SKU")),
            consultedAt=iso_date(bs.get("DATA CONSULTA")),
            commercialMode=row.get("MODALIDADE COMERCIAL"),
            commercialIdentityKey=row.get("CHAVE IDENTIDADE COMERCIAL"),
            validatedContents=split_codes(row.get("CONTEÚDOS VALIDADOS")),
            pendingVariationCodes=split_codes(row.get("PENDÊNCIAS DE CAMPOS")),
            values=values,
        )
        if not values:
            report["sku_without_values"].append([source_scode, pcode, product["name"]])
        product["skus"].append(sku)

    for pcode, source_scode in pending_bare:
        seq = 1
        while seq in sku_seq_used[pcode]:
            seq += 1
        sku_seq_used[pcode].add(seq)
        new_code = f"{pcode}.S{seq:04d}"
        for sku in products[pcode]["skus"]:
            if sku["sourceCode"] == source_scode:
                sku["skuCode"] = new_code

    for product in products.values():
        product["skus"].sort(key=lambda s: code_sort_key(s["skuCode"]))
        if product["sourceCode"]:
            report["code_normalized"].append(["PRODUTO", "sufixo com 3 dígitos", product["sourceCode"], product["productCode"]])
        for sku in product["skus"]:
            if sku["sourceCode"] == sku["skuCode"]:
                sku["sourceCode"] = None
            elif not sku["sourceCode"].startswith(product["sourceCode"] or product["productCode"]):
                report["code_normalized"].append(["SKU", "código sem o produto", sku["sourceCode"], sku["skuCode"]])
            else:
                report["code_normalized"].append(["SKU", "sufixo com 3 dígitos" if not product["sourceCode"] else "produto renumerado",
                                                  sku["sourceCode"], sku["skuCode"]])

    # Mesmo código de fabricante + mesma marca em produtos diferentes, sem vínculo de identidade declarado.
    by_reference = defaultdict(list)
    for product in products.values():
        for sku in product["skus"]:
            if sku["manufacturerSku"] and product["brand"]:
                by_reference[(key(product["brand"]), key(sku["manufacturerSku"]))].append((product, sku))
    for (_, reference), items in sorted(by_reference.items()):
        if len({p["productCode"] for p, _ in items}) < 2:
            continue
        if len({s["commercialIdentityKey"] for _, s in items}) == 1 and items[0][1]["commercialIdentityKey"]:
            continue
        for product, sku in items:
            report["possible_duplicates"].append([sku["manufacturerSku"], product["brand"], sku["skuCode"],
                                                  product["materialCode"], product["name"], sku["presentation"]])
    return products


def split_codes(text):
    if not text:
        return None
    codes = [c.strip() for c in re.split(r"[|;,]", text) if c.strip()]
    return codes or None


def build_value(item, material, catalog, sku_code, sku_source, report):
    vcode = item["codigo_variacao"]
    variations = catalog[material]["variations"]
    if vcode not in variations:
        raise SystemExit(f"{sku_code}: variação {vcode} não pertence ao material {material}.")
    definition = variations[vcode]["variation"]
    raw = item.get("valor")
    unit = clean_text(item.get("unidade"))
    numeric = None
    dims = None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        numeric = float(raw)
        display = format_number(raw)
        if definition.get("FORMATO") == "DIMENSOES":
            dims = [numeric]
    elif raw is None:
        display = None
    else:
        display = clean_text(raw)
        if definition.get("FORMATO") == "DIMENSOES" and display:
            dims, display = parse_dimensions(display)
            if dims is None:
                report["value_not_parsed"].append([sku_code, vcode, item.get("atributo"), raw])
            elif len(dims) == 1:
                numeric = dims[0]
    option_code = None
    options = variations[vcode]["options"]
    if options and display is not None:
        candidates = {option_keys(display), option_keys(display, unit)}
        if numeric is not None:
            candidates |= {option_keys(format_number(numeric)), option_keys(format_number(numeric), unit)}
        for o in options:
            keys = {option_keys(o["OPÇÃO"]), option_keys(o.get("SÍMBOLO / UNIDADE")),
                    option_keys(o["OPÇÃO"], o.get("SÍMBOLO / UNIDADE"))}
            if candidates & keys:
                option_code = o["Cod. OP"]
                break
        if option_code is None:
            report["value_without_option"].append([sku_code, vcode, definition["VARIAÇÃO"], display, unit,
                                                   " | ".join(o["OPÇÃO"] for o in options[:12])])
    if unit and definition.get("UNIDADES_ACEITAS") and unit not in definition["UNIDADES_ACEITAS"]:
        report["unit_not_accepted"].append([sku_code, vcode, unit, " | ".join(definition["UNIDADES_ACEITAS"])])
    return OrderedDict(
        variationCode=vcode,
        attribute=definition["VARIAÇÃO"],
        value=display,
        numericValue=numeric,
        dimensions=dims,
        unit=unit,
        optionCode=option_code,
        qualifier=clean_text(item.get("qualificador")),
        originalValue=clean_text(item.get("valor_original_documentado")),
        sourceUrl=clean_text(item.get("fonte")) or clean_text(sku_source),
    )


# --------------------------------------------------------------------------------------------
# Checagens finais (falham a execução se algo estiver errado)
# --------------------------------------------------------------------------------------------

def validate(products, catalog):
    sku_codes = Counter()
    gtins = Counter()
    for p in products.values():
        assert re.fullmatch(re.escape(p["materialCode"]) + r"\.P\d{4,5}", p["productCode"]), p["productCode"]
        assert p["materialCode"] in catalog, p["productCode"]
        assert p["name"], p["productCode"]
        assert p["skus"], p["productCode"]
        for s in p["skus"]:
            assert re.fullmatch(re.escape(p["productCode"]) + r"\.S\d{4}", s["skuCode"]), s["skuCode"]
            sku_codes[s["skuCode"]] += 1
            if s["gtin"]:
                gtins[s["gtin"]] += 1
            seen = set()
            for v in s["values"]:
                assert v["variationCode"] in catalog[p["materialCode"]]["variations"], (s["skuCode"], v)
                assert v["variationCode"] not in seen, (s["skuCode"], v["variationCode"])
                seen.add(v["variationCode"])
                if v["optionCode"]:
                    opts = {o["Cod. OP"] for o in catalog[p["materialCode"]]["variations"][v["variationCode"]]["options"]}
                    assert v["optionCode"] in opts, (s["skuCode"], v)
        doc = json.dumps(p, ensure_ascii=False)
        assert len(doc.encode()) < 1_000_000, p["productCode"]
    dup = [c for c, n in sku_codes.items() if n > 1]
    assert not dup, f"SKUs duplicados: {dup[:5]}"
    dup_gtin = [c for c, n in gtins.items() if n > 1]
    assert not dup_gtin, f"GTIN duplicado: {dup_gtin[:5]}"


def prune(obj):
    """Remove campos nulos/vazios (o backend serializa com non_null)."""
    if isinstance(obj, dict):
        return OrderedDict((k, prune(v)) for k, v in obj.items() if v is not None and v != [] and v != {})
    if isinstance(obj, list):
        return [prune(v) for v in obj]
    return obj


# --------------------------------------------------------------------------------------------
# Relatório
# --------------------------------------------------------------------------------------------

REPORT_SHEETS = OrderedDict([
    ("excluded_rows", ("Linhas excluídas", ["SKU", "Produto", "Motivo"])),
    ("catalog_new_material", ("Materiais novos", ["Cód. material", "Família", "Material", "Família nova?"])),
    ("catalog_material_renamed", ("Materiais renomeados", ["Cód. material", "Nome no banco", "Nome do cliente"])),
    ("catalog_family_renamed", ("Famílias renomeadas", ["Cód. material", "Cód. família", "Nome no banco", "Nome do cliente"])),
    ("catalog_material_not_in_contract", ("Materiais fora do contrato", ["Cód. material", "Família", "Material"])),
    ("catalog_new_variation", ("Variações novas", ["Cód. variação", "Material", "Atributo", "Formato"])),
    ("catalog_variation_renamed", ("Variações renomeadas", ["Cód. variação", "Nome no banco", "Nome do cliente", "Tipo no banco", "Tipo do cliente", "Opções no banco", "Leitura"])),
    ("catalog_variation_not_in_contract", ("Variações fora do contrato", ["Cód. variação", "Material", "Atributo", "Opções no banco"])),
    ("code_normalized", ("Códigos reestruturados", ["Nível", "Motivo", "Código do cliente", "Código no banco"])),
    ("possible_duplicates", ("Possíveis duplicidades", ["SKU fabricante", "Marca", "SKU", "Material", "Produto", "Apresentação"])),
    ("brand_variants", ("Grafias de marca", ["Campo", "Grafia adotada", "Grafias encontradas", "Ocorrências"])),
    ("unit_not_standard", ("Unidade não padronizada", ["SKU", "Unidade comercial informada"])),
    ("numeric_cells", ("Células numéricas", ["SKU", "Campo", "Valor lido", "Observação"])),
    ("gtin_invalid", ("GTIN inválido", ["SKU", "GTIN"])),
    ("value_not_parsed", ("Dimensões não lidas", ["SKU", "Cód. variação", "Atributo", "Valor"])),
    ("value_without_option", ("Valor sem opção", ["SKU", "Cód. variação", "Atributo", "Valor", "Unidade", "Opções do catálogo"])),
    ("unit_not_accepted", ("Unidade fora do contrato", ["SKU", "Cód. variação", "Unidade", "Unidades aceitas"])),
    ("sku_without_values", ("SKUs sem valores", ["SKU", "Produto", "Nome"])),
])


def client_actions(report):
    """Itens que dependem de decisão do cliente (só aparecem quando há ocorrências; exemplos vêm do próprio lote)."""
    n = {k: f"{len(v):,}".replace(",", ".") for k, v in report.items()}
    has = {k: len(v) for k, v in report.items()}
    reused_rows = [r for r in report["catalog_variation_renamed"] if r[-1].startswith("REVISAR")]

    def example(rows, render):
        return f" (ex.: {render(rows[0])})" if rows else ""

    items = [
        (len(reused_rows), f"{len(reused_rows)} códigos de variação foram reaproveitados com outro significado"
                           + example(reused_rows, lambda r: f"{r[0]}: '{r[1]}' → '{r[2]}'")
                           + ". Confirmar; o ideal é criar código novo, como já feito em outros casos."),
        (has.get("catalog_material_not_in_contract", 0) + has.get("catalog_variation_not_in_contract", 0),
         f"{n.get('catalog_material_not_in_contract', 0)} materiais e {n.get('catalog_variation_not_in_contract', 0)} variações do nosso "
         "catálogo não estão no contrato do cliente. Foram mantidos; confirmar se saíram ou foram renumerados."),
        (has.get("possible_duplicates", 0), f"{n.get('possible_duplicates', 0)} SKUs compartilham marca + código do fabricante com outro "
                                            "produto sem chave de identidade comercial"
                                            + example(report["possible_duplicates"], lambda r: f"{r[1]} {r[0]}")
                                            + ". Confirmar se é o mesmo item."),
        (has.get("unit_not_standard", 0), f"{n.get('unit_not_standard', 0)} SKUs com UNIDADE COMERCIAL fora do padrão"
                                          + example(report["unit_not_standard"], lambda r: f"'{r[1]}'") + "."),
        (has.get("value_without_option", 0), f"{n.get('value_without_option', 0)} valores documentados não existem nas opções do catálogo"
                                             + example(report["value_without_option"], lambda r: f"{r[2]} = {r[3]}{(' ' + r[4]) if r[4] else ''}")
                                             + ". Ampliar as opções ou tratar o campo como livre."),
        (has.get("sku_without_values", 0), f"{n.get('sku_without_values', 0)} SKUs aptos sem nenhum valor de variação "
                                           "(só identidade e apresentação)."),
        (has.get("code_normalized", 0),
         f"{n.get('code_normalized', 0)} códigos fora do padrão (sufixo com 3 dígitos, SKU sem o código do produto). "
         "Padronizar na origem; o de-para completo está na aba Códigos reestruturados."),
        (has.get("brand_variants", 0), f"{n.get('brand_variants', 0)} marcas/fabricantes com grafias diferentes"
                                       + example(report["brand_variants"], lambda r: r[2]) + "."),
    ]
    return [text for count, text in items if count]


def write_report(path, report, summary):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    head = Font(name="Arial", bold=True, color="FFFFFF")
    fill = PatternFill("solid", start_color="1F3A5F")
    body = Font(name="Arial", size=10)
    ws.append(["Indicador", "Valor"])
    for k, v in summary:
        ws.append([k, v])
    ws.append([])
    ws.append(["Aba do relatório", "Linhas"])
    sub_row = ws.max_row
    for code, (title, _) in REPORT_SHEETS.items():
        ws.append([title, len(report[code])])
    ws.append([])
    ws.append(["Pontos para confirmar com o cliente", ""])
    actions_row = ws.max_row
    for text in client_actions(report):
        ws.append([text, ""])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
    for row in ws.iter_rows():
        for c in row:
            c.font = body
            c.alignment = Alignment(vertical="top", wrap_text=c.row > actions_row)
            if isinstance(c.value, int):
                c.number_format = "#,##0"
    for r in (1, sub_row, actions_row):
        for c in ws[r]:
            c.font, c.fill = head, fill
    for r in range(actions_row + 1, ws.max_row + 1):
        ws.row_dimensions[r].height = 30
    ws.column_dimensions["A"].width = 68
    ws.column_dimensions["B"].width = 70
    for code, (title, columns) in REPORT_SHEETS.items():
        sh = wb.create_sheet(title[:31])
        sh.append(columns)
        for row in report[code]:
            sh.append([str(v) if isinstance(v, (list, dict)) else v for v in row])
        for c in sh[1]:
            c.font, c.fill = head, fill
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.font = body
                c.alignment = Alignment(vertical="top", wrap_text=False)
        for i, col in enumerate(columns, start=1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in report[code][:500] if i - 1 < len(r)])
            sh.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 70)
        sh.freeze_panes = "A2"
        if report[code]:
            sh.auto_filter.ref = sh.dimensions
    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)


# --------------------------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx")
    parser.add_argument("--repo", default=".")
    parser.add_argument("--report", default=None)
    args = parser.parse_args()
    repo = Path(args.repo)
    source = Path(args.xlsx)

    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = CalamineWorkbook.from_path(str(source))
    manifest_sheet = {r["CAMPO"]: r["VALOR"] for r in read_sheet(workbook, "00_MANIFESTO")}
    source_date = (manifest_sheet.get("DATA") or dt.date.today().isoformat())[:10]
    batch_id = f"PRICO_{source_date.replace('-', '')}_{sha[:8]}"

    report = defaultdict(list)

    v1 = read_bundled_catalog(repo)
    contracts = read_sheet(workbook, "03_CONTRATO_VARIACOES")
    coverage = read_sheet(workbook, "28_COBERTURA_MATERIAIS")
    catalog_rows, catalog = build_catalog(v1, contracts, coverage, report)

    apto = read_sheet(workbook, "37_CATALOGO_APTO")
    products = build_products(apto, read_sheet(workbook, "10_PRODUTOS"), read_sheet(workbook, "11_SKUS"),
                              catalog, report, batch_id)
    validate(products, catalog)

    # Catálogo v2 no mesmo formato do arquivo atual (linhas por opção).
    ndjson = "\n".join(json.dumps(r, ensure_ascii=False) for r in catalog_rows) + "\n"
    buf = io.BytesIO()
    with gzip.GzipFile(fileobj=buf, mode="wb", mtime=0) as gz:
        gz.write(ndjson.encode("utf-8"))
    (repo / CATALOG_RESOURCE).write_bytes(base64.encodebytes(buf.getvalue()))

    lines = [json.dumps(prune(p), ensure_ascii=False) for p in products.values()]
    buf = io.BytesIO()
    with gzip.GzipFile(fileobj=buf, mode="wb", mtime=0) as gz:
        gz.write(("\n".join(lines) + "\n").encode("utf-8"))
    (repo / PRODUCTS_RESOURCE).write_bytes(buf.getvalue())
    products_sha = hashlib.sha256(buf.getvalue()).hexdigest()

    sku_total = sum(len(p["skus"]) for p in products.values())
    value_total = sum(len(s["values"]) for p in products.values() for s in p["skus"])
    observation = next((r.get("OBSERVAÇÃO") for r in apto if r.get("OBSERVAÇÃO")), None)
    manifest = OrderedDict(
        batchId=batch_id,
        sourceFile=source.name,
        sourceSha256=sha,
        sourceDate=source_date,
        productsSha256=products_sha,
        generatedBy="tools/prico/consolidar_prico.py",
        rule="Somente 37_CATALOGO_APTO com PROTOCOLO PARA INSERÇÃO = OK e controles do registro conferidos.",
        observation=observation,
        counts=OrderedDict(
            materials=len(catalog),
            variations=sum(len(m["variations"]) for m in catalog.values()),
            options=sum(len(v["options"]) for m in catalog.values() for v in m["variations"].values()),
            products=len(products),
            skus=sku_total,
            values=value_total,
            excludedRows=len(report["excluded_rows"]),
        ),
    )
    (repo / MANIFEST_RESOURCE).write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    summary = [
        ("Lote", batch_id),
        ("Arquivo de origem", source.name),
        ("SHA-256 da origem", sha),
        ("Linhas em 37_CATALOGO_APTO", len(apto)),
        ("Linhas excluídas", len(report["excluded_rows"])),
        ("Produtos consolidados", len(products)),
        ("SKUs consolidados", sku_total),
        ("Valores de variação", value_total),
        ("Valores ligados a uma opção do catálogo", sum(1 for p in products.values() for s in p["skus"] for v in s["values"] if v["optionCode"])),
        ("Materiais no catálogo v2", manifest["counts"]["materials"]),
        ("Variações no catálogo v2", manifest["counts"]["variations"]),
        ("Opções no catálogo v2", manifest["counts"]["options"]),
        ("Produtos com código reajustado (P001 → P0001)", sum(1 for p in products.values() if p["sourceCode"])),
        ("SKUs com código reajustado (S001 → S0001, S00001 → hierárquico)", sum(1 for p in products.values() for s in p["skus"] if s["sourceCode"])),
    ]
    report_path = Path(args.report) if args.report else repo / f"relatorio_revisao_{batch_id}.xlsx"
    write_report(report_path, report, summary)
    for k, v in summary:
        print(f"{k:45s} {v}")
    for code, (title, _) in REPORT_SHEETS.items():
        print(f"  {title:35s} {len(report[code])}")
    print(f"Relatório: {report_path}")


if __name__ == "__main__":
    sys.exit(main())
